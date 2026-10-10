from fastapi import APIRouter, UploadFile, File, HTTPException
from typing import List, Optional
from app.rag.pipeline import index_document


router = APIRouter(
    prefix="/api/documents",
    tags=["AI Documents"],
)


# ---------------------------------------------------------
# Supported file types
# ---------------------------------------------------------

ALLOWED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".csv",
    ".png",
    ".jpg",
    ".jpeg",
}


from pydantic import BaseModel
from pathlib import Path
import io

class ProcessFileRequest(BaseModel):
    document_id: str
    file_path: str
    filename: Optional[str] = None
    user_id: Optional[str] = None


from app.document_processing.pdf import is_valid_page_text, clean_page_text


def _extract_from_pdf_bytes(content: bytes) -> str:
    pages_text = []

    # Strategy 1: PyMuPDF (fitz)
    try:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
        doc = fitz.open(stream=content, filetype="pdf")
        for page in doc:
            raw = page.get_text("text") or ""
            cleaned = clean_page_text(raw)
            if is_valid_page_text(cleaned):
                pages_text.append(cleaned)
            else:
                pages_text.append("")
        doc.close()
    except Exception as fitz_err:
        print(f"[PDF Bytes] PyMuPDF extraction warning: {fitz_err}")
        pages_text = []

    # Strategy 2: pypdf fallback
    if not pages_text or sum(len(p) for p in pages_text) < 50:
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            pages_text = []
            for page in reader.pages:
                raw = page.extract_text() or ""
                cleaned = clean_page_text(raw)
                if is_valid_page_text(cleaned):
                    pages_text.append(cleaned)
                else:
                    pages_text.append("")
        except Exception as pypdf_err:
            print(f"[PDF Bytes] pypdf extraction warning: {pypdf_err}")

    combined = "\f".join(pages_text)
    if sum(len(p) for p in pages_text) >= 50:
        return combined

    # Strategy 3: OCR Fallback if normal extraction yielded no usable text
    print("[OCR] Triggered OCR fallback for PDF without readable text layer")
    from app.document_processing.ocr import find_tesseract_binary, get_ocr_status
    ocr_stat = get_ocr_status()
    if not ocr_stat["pytesseract_python"] or not ocr_stat["tesseract_binary"]:
        print("[OCR] OCR_UNAVAILABLE: Scanned PDF detected, but OCR engine is unavailable.")
        return "OCR_UNAVAILABLE"

    try:
        import importlib
        pytesseract = importlib.import_module("pytesseract")
        pil_img_mod = importlib.import_module("PIL.Image")
        Image = getattr(pil_img_mod, "Image", pil_img_mod)
        t_bin = find_tesseract_binary()
        if t_bin:
            pytesseract.pytesseract.tesseract_cmd = t_bin

        ocr_pages = []
        try:
            try:
                import pymupdf as fitz
            except ImportError:
                import fitz
            doc = fitz.open(stream=content, filetype="pdf")
            for i, page in enumerate(doc):
                if i >= 25:
                    break
                pix = page.get_pixmap(dpi=200)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                txt = clean_page_text(pytesseract.image_to_string(img).strip())
                if txt:
                    ocr_pages.append(txt)
            doc.close()
        except Exception as fitz_ocr_err:
            print(f"[OCR] PyMuPDF pixmap OCR notice: {fitz_ocr_err}")

        if ocr_pages:
            return "\f".join(ocr_pages)

        # Fallback to pdf2image if available
        try:
            pdf2image = importlib.import_module("pdf2image")
            convert_from_bytes = getattr(pdf2image, "convert_from_bytes", None)
            if convert_from_bytes:
                images = convert_from_bytes(content, first_page=1, last_page=25)
                ocr_pages = [clean_page_text(pytesseract.image_to_string(img).strip()) for img in images]
                return "\f".join(ocr_pages)
        except Exception as p2i_err:
            print(f"[OCR] pdf2image notice: {p2i_err}")

    except Exception as ocr_err:
        print(f"[OCR] OCR fallback warning: {ocr_err}")

    return combined if combined.strip() else "OCR_UNAVAILABLE"


def _extract_from_docx_bytes(content: bytes) -> str:
    paragraphs = []
    # Strategy 1: python-docx if installed
    try:
        import docx
        doc = docx.Document(io.BytesIO(content))
        for p_item in doc.paragraphs:
            if p_item.text.strip():
                paragraphs.append(p_item.text.strip())
        for tab in doc.tables:
            for row in tab.rows:
                row_txt = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_txt:
                    paragraphs.append(row_txt)
        if paragraphs:
            return "\n\n".join(paragraphs)
    except Exception as docx_err:
        print(f"[DOCX Bytes] python-docx notice: {docx_err}; using zipfile XML")

    # Strategy 2: Standard library zipfile + WordML XML
    try:
        import zipfile
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if "word/document.xml" in z.namelist():
                tree = ET.fromstring(z.read("word/document.xml"))
                ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                body = tree.find("w:body", ns)
                if body is not None:
                    for child in body:
                        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
                        if tag == "p":
                            texts = [node.text for node in child.findall(".//w:t", ns) if node.text]
                            p_txt = "".join(texts).strip()
                            if p_txt:
                                paragraphs.append(p_txt)
                        elif tag == "tbl":
                            for tr in child.findall("w:tr", ns):
                                cells = []
                                for tc in tr.findall("w:tc", ns):
                                    c_texts = [node.text for node in tc.findall(".//w:t", ns) if node.text]
                                    cells.append("".join(c_texts).strip())
                                if any(cells):
                                    paragraphs.append(" | ".join(cells))
    except Exception as zip_err:
        print(f"[DOCX Bytes] zipfile WordML extraction error: {zip_err}")

    return "\n\n".join(paragraphs)


@router.post("/process")
async def process_document(file: UploadFile = File(...)):

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="No filename provided."
        )

    filename = file.filename.lower()

    if not any(filename.endswith(ext) for ext in ALLOWED_EXTENSIONS):
        raise HTTPException(
            status_code=400,
            detail="Unsupported document type."
        )

    content = await file.read()
    text = ""
    if filename.endswith(".pdf"):
        text = _extract_from_pdf_bytes(content)
    elif filename.endswith((".docx", ".doc")):
        text = _extract_from_docx_bytes(content)
    elif filename.endswith((".txt", ".csv", ".json", ".md")):
        text = content.decode("utf-8", errors="ignore")
    if text == "OCR_UNAVAILABLE":
        raise HTTPException(status_code=422, detail="OCR_UNAVAILABLE")
    if not text.strip():
        raise HTTPException(status_code=422, detail="No readable text was extracted from the document")
    document_id = file.headers.get("x-document-id") or file.filename
    user_id = file.headers.get("x-user-id") or ""
    chunks = index_document(document_id, text, file.filename, user_id=user_id)
    return {
        "success": True,
        "filename": file.filename,
        "content_type": file.content_type,
        "status": "indexed",
        "chunks": chunks,
        "document_id": document_id,
    }


@router.post("/process-file")
async def process_file_path(request: ProcessFileRequest):
    """
    Process and index an uploaded document using its path on the local server filesystem.
    """
    p = Path(request.file_path)
    if not p.exists():
        raise HTTPException(status_code=404, detail=f"File not found on server at {request.file_path}")

    filename = request.filename or p.name
    ext = p.suffix.lower()

    content = p.read_bytes()
    text = ""
    if ext == ".pdf":
        text = _extract_from_pdf_bytes(content)
    elif ext in {".docx", ".doc"}:
        text = _extract_from_docx_bytes(content)
    elif ext in {".txt", ".csv", ".json", ".md", ".log"}:
        text = content.decode("utf-8", errors="ignore")
    else:
        text = content.decode("utf-8", errors="ignore")

    if text == "OCR_UNAVAILABLE":
        raise HTTPException(status_code=422, detail="OCR_UNAVAILABLE")
    if not text.strip():
        raise HTTPException(status_code=422, detail="No readable text was extracted from the file")

    chunks = index_document(
        document_id=request.document_id,
        text=text,
        filename=filename,
        user_id=request.user_id or "",
    )

    return {
        "success": True,
        "filename": filename,
        "status": "indexed",
        "chunks": chunks,
        "document_id": request.document_id,
    }



# ---------------------------------------------------------
# Supported File Types
# ---------------------------------------------------------

@router.get("/supported-types")
async def supported_types():

    return {
        "success": True,
        "supported_types": sorted(ALLOWED_EXTENSIONS)
    }