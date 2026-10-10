#!/usr/bin/env python3
"""
DOCUMIND-AI: Centralized High-Fidelity Document Loader Service
Supports:
- PDF (native PyMuPDF page-by-page text & table extraction, OCR fallback)
- DOCX (structured paragraphs & tables)
- TXT, CSV, MD, JSON
- Images (JPG, JPEG, PNG via OCR)

Returns normalized JSON format:
{
  "documentId": "...",
  "fileName": "...",
  "fileType": "pdf",
  "totalPages": 23,
  "pages": [
    {
      "pageNumber": 1,
      "text": "...",
      "tables": [...],
      "section": "Introduction",
      "isReference": false,
      "extractionMethod": "native",
      "ocrUsed": false,
      "ocrConfidence": 1.0,
      "qualityScore": 0.98
    }
  ],
  "fullText": "..."
}
"""

import sys
import os
import re
import json
import io
import csv
from pathlib import Path
import warnings

# Ensure UTF-8 output on Windows stdout/stderr for robust Node-Python IPC
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="strict")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

warnings.filterwarnings("ignore", category=UserWarning)

# Bytecode and xref regexes for text quality validation
# Detects raw PDF file structure (when a binary PDF file was opened as text),
# while never falsely matching standard human English words like "trailer", "xref", or "/Font".
INVALID_PDF_PATTERNS = [
    re.compile(r"\bendstream\b", re.IGNORECASE),
    re.compile(r"\bendobj\b", re.IGNORECASE),
    re.compile(r"\b\d+\s+\d+\s+obj\b", re.IGNORECASE),
    re.compile(r"/CIDInit\b", re.IGNORECASE),
    re.compile(r"/ProcSet\s*\[", re.IGNORECASE),
    re.compile(r"/Font\s*<<", re.IGNORECASE),
    re.compile(r"/CMapName\b", re.IGNORECASE),
    re.compile(r"/Adobe-Identity-UCS\b", re.IGNORECASE),
    re.compile(r"/MediaBox\s*\[", re.IGNORECASE),
    re.compile(r"/Contents\s+\d+\s+\d+\s+R\b", re.IGNORECASE),
    re.compile(r"/BaseFont\s*/", re.IGNORECASE),
    re.compile(r"/CIDFont\b", re.IGNORECASE),
    re.compile(r"/CIDSystemInfo\s*<<", re.IGNORECASE),
    re.compile(r"/FlateDecode\b", re.IGNORECASE),
    re.compile(r"^%PDF-\d\.\d", re.IGNORECASE),
    re.compile(r"\bstartxref\s*\n", re.IGNORECASE),
    re.compile(r"\btrailer\s*<<", re.IGNORECASE),
    re.compile(r"\bxref\s*\n\d+", re.IGNORECASE),
    re.compile(r"\b\d{10}\s+\d{5}\s+[nf]\b"),
    re.compile(r"\b00000\s+65535\s+f\b"),
]

# Section patterns for academic & enterprise document structure detection
SECTION_PATTERNS = [
    ("Abstract", re.compile(r"^(?:(?:\d+[\.\s]+)?abstract)\b", re.IGNORECASE)),
    ("Introduction", re.compile(r"^(?:(?:\d+[\.\s]+)?(?:introduction|background))\b", re.IGNORECASE)),
    ("Motivation", re.compile(r"^(?:(?:\d+[\.\s]+)?(?:research\s+motivation|motivation))\b", re.IGNORECASE)),
    ("Problem Statement", re.compile(r"^(?:(?:\d+[\.\s]+)?(?:problem\s+statement|research\s+problem|deficiencies))\b", re.IGNORECASE)),
    ("Literature Review", re.compile(r"^(?:(?:\d+[\.\s]+)?(?:literature\s+review|related\s+work))\b", re.IGNORECASE)),
    ("Research Gap", re.compile(r"^(?:(?:\d+[\.\s]+)?(?:research\s+gap|limitations\s+of\s+prior\s+work))\b", re.IGNORECASE)),
    ("Objectives", re.compile(r"^(?:(?:\d+[\.\s]+)?(?:research\s+objectives?|objectives?|aims?))\b", re.IGNORECASE)),
    ("Methodology", re.compile(r"^(?:(?:\d+[\.\s]+)?(?:methodology|proposed\s+framework|proposed\s+approach|system\s+model|vla\s+framework))\b", re.IGNORECASE)),
    ("System Architecture", re.compile(r"^(?:(?:\d+[\.\s]+)?(?:system\s+architecture|architecture|framework\s+architecture))\b", re.IGNORECASE)),
    ("Algorithms", re.compile(r"^(?:(?:\d+[\.\s]+)?(?:algorithms?|pseudocode))\b", re.IGNORECASE)),
    ("Evaluation", re.compile(r"^(?:(?:\d+[\.\s]+)?(?:evaluation|experiments?|benchmarks?|datasets?|experimental\s+results?))\b", re.IGNORECASE)),
    ("Expected Contributions", re.compile(r"^(?:(?:\d+[\.\s]+)?(?:expected\s+contributions?|contributions?))\b", re.IGNORECASE)),
    ("Roadmap & Timeline", re.compile(r"^(?:(?:\d+[\.\s]+)?(?:timeline|work\s+plan|implementation\s+plan|schedule|milestones))\b", re.IGNORECASE)),
    ("Conclusion", re.compile(r"^(?:(?:\d+[\.\s]+)?(?:conclusion|concluding\s+remarks|summary\s+and\s+future\s+work))\b", re.IGNORECASE)),
    ("References", re.compile(r"^(?:(?:\d+[\.\s]+)?(?:references|bibliography|works\s+cited))\b", re.IGNORECASE)),
]


def assess_text_quality(text: str, file_type: str = "pdf") -> dict:
    """
    STAGE 5: Text Quality Validation
    Returns: { 'valid': bool, 'score': float, 'reason': str }
    """
    if not text or not isinstance(text, str):
        return {"valid": False, "score": 0.0, "reason": "Empty or non-string text"}

    trimmed = text.strip()
    if len(trimmed) < 15:
        return {"valid": False, "score": 0.0, "reason": "Text is too short (< 15 characters)"}

    # 1. Reject PDF bytecode / xref tables (only for raw PDF header/bytecode dumps)
    if file_type == "pdf" or "%PDF-" in trimmed[:20]:
        for pat in INVALID_PDF_PATTERNS:
            if pat.search(trimmed):
                return {"valid": False, "score": 0.0, "reason": f"PDF internal object/xref artifact detected: {pat.pattern}"}

    # 2. Reject excessive null bytes
    null_count = trimmed.count("\x00")
    if null_count > 0 and (null_count / len(trimmed)) > 0.02:
        return {"valid": False, "score": 0.0, "reason": "Excessive null bytes"}

    # 3. Normalize table formatting before ratio checks
    content_no_grid = trimmed.replace("|", " ")
    is_tabular = "|" in trimmed or file_type in ("csv", "tsv")
    is_csv = file_type in ("csv", "tsv")
    is_docx = file_type in ("docx", "doc")
    has_multilingual = any(ord(c) > 0x0400 for c in trimmed)

    # 4. Calculate alphanumeric ratio
    letters = sum(1 for c in content_no_grid if c.isalnum() or c.isspace())
    alpha_ratio = letters / len(content_no_grid) if len(content_no_grid) > 0 else 0
    min_alpha = 0.15 if is_csv else (0.20 if (is_tabular or is_docx or has_multilingual) else 0.25)
    if alpha_ratio < min_alpha:
        return {"valid": False, "score": round(alpha_ratio, 2), "reason": f"Low alphanumeric ratio ({alpha_ratio:.2f})"}

    # 5. Check for single-character glyph token spam (skip for tables, CSV, DOCX, multilingual)
    # Strip bullet symbols and list markers so structured bullet/numbered lists are not treated as glyph noise
    content_no_bullets = re.sub(r"[•▪▫◦–—·\*\-│┌┬┐└┴┘├┼┤═║]", " ", content_no_grid)
    tokens = [t for t in content_no_bullets.split() if not re.match(r"^\d+[\.\)]?$", t)]
    if len(tokens) >= 10 and not has_multilingual and not is_tabular and not is_docx and not is_csv:
        short_tokens = sum(1 for t in tokens if len(t) <= 2)
        short_ratio = short_tokens / len(tokens) if tokens else 0
        if short_ratio > 0.75:
            return {"valid": False, "score": 0.2, "reason": f"Garbled glyph token noise: {short_ratio:.1%}"}

    words = [t for t in tokens if len(t) >= 2 and any(c.isalnum() for c in t)]
    word_density = len(words) / len(tokens) if tokens else 0
    score = round(min(1.0, max(0.4, word_density * 0.7 + alpha_ratio * 0.3)), 2)

    return {"valid": True, "score": score, "reason": "Clean readable text"}


def clean_text_content(text: str) -> str:
    """
    STAGE 8: Text Cleaning
    Cleans:
    - repeated headers & footers
    - broken line breaks & hyphenated word splits
    - null bytes & replacement characters
    - malformed whitespace
    Preserves:
    - paragraph breaks, headings, lists, table structures
    """
    if not text:
        return ""

    # Strip replacement character & null bytes
    cleaned = text.replace("\ufffd", " ").replace("\x00", "")
    # Strip HTML / SVG artifact tags
    cleaned = re.sub(r"<svg[\s\S]*?</svg>", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"<[^>]+>", " ", cleaned)

    # Reconnect hyphenated line splits: e.g. "Generat-\nion" -> "Generation"
    cleaned = re.sub(r"(\w+)-\s*\n\s*(\w+)", r"\1\2", cleaned)

    lines = []
    for line in cleaned.splitlines():
        trimmed = line.strip()
        if not trimmed:
            continue
        # Strip common repeating running headers & footers
        if re.match(r"^MERA-VLA\s*—.*Page\s+\d+\s+of\s+\d+", trimmed, re.IGNORECASE):
            continue
        if re.match(r"^Page\s+\d+\s+of\s+\d+$", trimmed, re.IGNORECASE):
            continue
        if re.match(r"^\d+\s*/\s*\d+$", trimmed):
            continue
        lines.append(trimmed)

    # Intelligently join paragraphs
    joined_paragraphs = []
    current_p = []

    for line in lines:
        if not current_p:
            current_p.append(line)
        else:
            prev = current_p[-1]
            # Paragraph or section heading boundary
            if prev.endswith((".", ":", "!", "?", "•", "—")) or line.startswith(("#", "###", "-", "*", "1.", "2.", "3.", "Table", "Figure")):
                joined_paragraphs.append(" ".join(current_p))
                current_p = [line]
            else:
                current_p.append(line)

    if current_p:
        joined_paragraphs.append(" ".join(current_p))

    return "\n\n".join(joined_paragraphs)


def detect_section_heading(text: str, current_section: str) -> tuple:
    """
    STAGE 7: Document Structure Detection
    Returns: (section_name: str, is_reference: bool)
    """
    lines = text.splitlines()[:6]
    for line in lines:
        trimmed = line.strip()
        for sec_name, pat in SECTION_PATTERNS:
            if pat.search(trimmed):
                is_ref = (sec_name == "References")
                return sec_name, is_ref
    return current_section, (current_section == "References")


def render_and_ocr_page(page, page_num: int) -> tuple:
    """
    STAGE 4: Page-by-Page OCR Fallback
    Renders PDF page to image at 200 DPI and extracts text via Tesseract.
    Returns: (ocr_text: str, ocr_confidence: float)
    """
    try:
        from PIL import Image
        import pytesseract

        # Check common Tesseract Windows installation paths
        for t_path in [
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
            os.path.expanduser(r"~\AppData\Local\Programs\Tesseract-OCR\tesseract.exe"),
        ]:
            if os.path.exists(t_path):
                pytesseract.pytesseract.tesseract_cmd = t_path
                break

        pix = page.get_pixmap(dpi=200)
        img = Image.open(io.BytesIO(pix.tobytes("png")))
        data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
        
        words = []
        confidences = []
        for i, word in enumerate(data.get("text", [])):
            w = word.strip()
            if w:
                words.append(w)
                try:
                    c = float(data["conf"][i])
                    if c >= 0:
                        confidences.append(c)
                except (ValueError, KeyError):
                    pass

        ocr_text = " ".join(words)
        avg_conf = (sum(confidences) / len(confidences) / 100.0) if confidences else 0.8
        return ocr_text, round(avg_conf, 2)
    except Exception as e:
        sys.stderr.write(f"[OCR] Page {page_num} OCR notice: {e}\n")
        return "", 0.0


def find_matching_caption_and_context(blocks, target_bbox, obj_prefix="tab"):
    """
    Finds the closest matching caption and nearby context for a table or figure.
    Returns: (clean_id, caption, nearby_context)
    """
    best_caption = ""
    best_id = ""
    min_dist = 999999.0
    pattern = re.compile(
        rf"^\s*(?:{obj_prefix}(?:ure)?|tab(?:le)?)\s*(\d+[a-zA-Z]?|\b[A-Z0-9.-]+)[:.\s]+(.*)",
        re.IGNORECASE,
    )

    t_y0 = target_bbox[1] if target_bbox and len(target_bbox) >= 4 else 0
    t_y1 = target_bbox[3] if target_bbox and len(target_bbox) >= 4 else 0

    nearby_lines = []

    for b in blocks:
        text = b[4].strip()
        m = pattern.search(text)
        if m:
            b_y0 = b[1]
            b_y1 = b[3]
            dist = min(abs(b_y0 - t_y1), abs(t_y0 - b_y1))
            if dist < min_dist:
                min_dist = dist
                best_id = m.group(1).lower()
                best_caption = text.replace("\n", " ")
        elif best_id and re.search(rf"\b(?:figure|fig\.|table)\s+{re.escape(best_id)}\b", text, re.IGNORECASE):
            nearby_lines.append(text.replace("\n", " "))

    nearby_context = " ".join(nearby_lines[:2])
    return best_id, best_caption, nearby_context


def extract_tables_from_page(page, page_num: int, current_section: str = "General", document_id: str = "", blocks: list = None) -> tuple:
    """
    STAGE 6: Layout-Aware Table Extraction
    Extracts structured tables with column headers, rows, cells, and dual representations
    (JSON structured data + searchable key-value format + Markdown table).
    Returns: (tables_list, evidence_objects_list)
    """
    tables = []
    evidence_objects = []
    if blocks is None:
        try:
            blocks = page.get_text("blocks")
        except Exception:
            blocks = []

    try:
        tabs = list(page.find_tables())
        for t_idx, tab in enumerate(tabs, start=1):
            extracted = tab.extract()
            if not extracted or len(extracted) < 2:
                continue

            bbox = list(tab.bbox) if hasattr(tab, "bbox") else []
            tab_num, caption, nearby_context = find_matching_caption_and_context(blocks, bbox, "tab")
            table_id = f"table_{tab_num}" if tab_num else f"table_{page_num}_{t_idx}"
            full_caption = caption or f"Table {t_idx} on Page {page_num}"

            # Extract clean headers and rows
            raw_headers = [str(col or "").strip() for col in extracted[0]]
            headers = [h if h else f"Col_{c_idx}" for c_idx, h in enumerate(raw_headers, start=1)]

            rows = []
            for row in extracted[1:]:
                clean_row = [str(cell or "").strip() for cell in row]
                rows.append(clean_row)

            # 1. Format Markdown representation
            header_line = "| " + " | ".join(headers) + " |"
            sep_line = "| " + " | ".join(["---"] * len(headers)) + " |"
            row_lines = ["| " + " | ".join(r) + " |" for r in rows]
            markdown_table = "\n".join([header_line, sep_line] + row_lines)

            # 2. Format Searchable Key-Value representation (optimized for dense & lexical retrieval)
            kv_row_lines = []
            for r_idx, r in enumerate(rows, start=1):
                pairs = []
                for c_idx, col_name in enumerate(headers):
                    val = r[c_idx] if c_idx < len(r) else ""
                    pairs.append(f"{col_name} = {val}")
                kv_row_lines.append(f"Row {r_idx}: {'; '.join(pairs)}")

            searchable_content = (
                f"[{full_caption} | Page {page_num} | Section: {current_section}]\n"
                f"Columns: {', '.join(headers)}\n\n"
                f"{'; '.join(kv_row_lines[:15])}\n\n"
                f"{markdown_table}"
            )

            structured_data = {
                "tableId": table_id,
                "page": page_num,
                "caption": full_caption,
                "columns": headers,
                "rows": rows,
                "rowCount": len(rows),
                "columnCount": len(headers),
                "bbox": bbox,
            }

            table_record = {
                "pageNumber": page_num,
                "tableId": table_id,
                "caption": full_caption,
                "headers": headers,
                "rows": rows,
                "rowCount": len(rows),
                "rawText": searchable_content,
                "markdown": markdown_table,
                "sourceBBox": bbox,
                "structuredData": structured_data,
            }
            tables.append(table_record)

            # Build EvidenceObject
            evidence_id = f"evidence_{document_id or 'doc'}_p{page_num}_{table_id}"
            evidence_objects.append({
                "id": evidence_id,
                "documentId": str(document_id),
                "pageNumber": page_num,
                "objectType": "table",
                "section": current_section,
                "objectId": table_id,
                "caption": full_caption,
                "content": searchable_content,
                "structuredData": structured_data,
                "bbox": bbox,
                "sourceMethod": "table_parser",
                "confidence": 0.98,
                "metadata": {
                    "columns": headers,
                    "rowCount": len(rows),
                    "columnCount": len(headers),
                    "nearbyContext": nearby_context,
                },
            })

            sys.stderr.write(f"[TABLE] document={document_id} page={page_num} table={table_id} rows={len(rows)} columns={len(headers)} status=success\n")

    except Exception as e:
        sys.stderr.write(f"[TABLE] Page {page_num} table extraction notice: {e}\n")

    return tables, evidence_objects


def extract_figures_from_page(page, page_num: int, current_section: str = "General", document_id: str = "", blocks: list = None, doc = None) -> tuple:
    """
    STAGE 6.5: Layout-Aware Figure & Image Extraction
    Extracts figures, diagrams, charts, and visual elements with OCR text,
    observable visual description, caption, and nearby contextual text.
    Returns: (figures_list, evidence_objects_list)
    """
    figures = []
    evidence_objects = []
    if blocks is None:
        try:
            blocks = page.get_text("blocks")
        except Exception:
            blocks = []

    try:
        imgs = page.get_images(full=True)
        seen_fig_ids = set()

        for img_idx, img in enumerate(imgs, start=1):
            xref = img[0]
            width = img[2]
            height = img[3]

            # Filter out tiny icon decorations, separators, or bullets (<50x50 px)
            if width < 50 or height < 50:
                continue

            rects = page.get_image_rects(xref)
            bbox = list(rects[0]) if rects else []

            fig_num, caption, nearby_context = find_matching_caption_and_context(blocks, bbox, "fig")
            fig_id = f"figure_{fig_num}" if fig_num else f"figure_{page_num}_{img_idx}"
            if fig_id in seen_fig_ids:
                fig_id = f"{fig_id}_{img_idx}"
            seen_fig_ids.add(fig_id)

            full_caption = caption or f"Figure {img_idx} on Page {page_num}"

            # Attempt OCR on figure image if pytesseract is available, and save image to disk
            ocr_text = ""
            img_path = None
            try:
                from PIL import Image
                import io
                if doc is not None:
                    pix = pymupdf.Pixmap(doc, xref)
                    if pix.n >= 5: # CMYK to RGB
                        pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                    img_data = pix.tobytes("png")
                    
                    # Save image to disk for multimodal vision analysis
                    script_root = Path(__file__).resolve().parent.parent
                    out_dir = script_root / "uploads" / "extracted_images" / str(document_id or "doc")
                    out_dir.mkdir(parents=True, exist_ok=True)
                    target_file = out_dir / f"pdf_p{page_num}_{fig_id}.png"
                    with open(target_file, "wb") as f_out:
                        f_out.write(img_data)
                    img_path = str(target_file.resolve())

                    try:
                        import pytesseract
                        pil_img = Image.open(io.BytesIO(img_data))
                        ocr_res = pytesseract.image_to_string(pil_img).strip()
                        if len(ocr_res) > 3:
                            ocr_text = re.sub(r"\s+", " ", ocr_res)
                    except Exception:
                        pass
            except Exception as save_err:
                sys.stderr.write(f"[FIGURE_SAVE_ERR] Page {page_num} fig {fig_id}: {save_err}\n")

            # Distinguish figure type (architecture diagram, workflow, flowchart, chart, graph)
            fig_type = "figure"
            lower_cap = full_caption.lower()
            if any(k in lower_cap for k in ["chart", "plot", "graph", "histogram", "curve"]):
                fig_type = "chart"
            elif any(k in lower_cap for k in ["architecture", "workflow", "pipeline", "diagram", "schematic"]):
                fig_type = "diagram"
            elif any(k in lower_cap for k in ["screenshot", "ui", "interface"]):
                fig_type = "screenshot"

            # Generate strictly observable visual description (never invent unread values)
            visual_desc_parts = [
                f"Visual {fig_type} ({width}x{height} px)",
                f"describing: {full_caption.replace('Figure ', '').replace('Fig. ', '')}",
            ]
            if nearby_context:
                visual_desc_parts.append(f"Operational context: {nearby_context}")
            visual_description = ". ".join(visual_desc_parts)

            exact_obj_id = f"Figure {fig_num}" if fig_num else fig_id

            # Searchable representation
            searchable_content = (
                f"[{full_caption} | Page {page_num} | Section: {current_section}]\n"
                f"Type: {fig_type.capitalize()}\n"
                f"Caption: {full_caption}\n"
                f"Visual Description: {visual_description}\n"
            )
            if ocr_text:
                searchable_content += f"OCR Embedded Text: {ocr_text}\n"
            if nearby_context:
                searchable_content += f"Nearby Text: {nearby_context}\n"

            structured_data = {
                "figureId": fig_id,
                "objectId": exact_obj_id,
                "figureNumber": str(fig_num) if fig_num else None,
                "page": page_num,
                "caption": full_caption,
                "figureType": fig_type,
                "ocrText": ocr_text,
                "visualDescription": visual_description,
                "nearbyContext": nearby_context,
                "imagePath": img_path,
                "width": width,
                "height": height,
                "bbox": bbox,
            }

            figure_record = {
                "pageNumber": page_num,
                "figureId": fig_id,
                "objectId": exact_obj_id,
                "figureNumber": str(fig_num) if fig_num else None,
                "caption": full_caption,
                "figureType": fig_type,
                "ocrText": ocr_text,
                "visualDescription": visual_description,
                "nearbyContext": nearby_context,
                "imagePath": img_path,
                "rawText": searchable_content,
                "sourceBBox": bbox,
                "structuredData": structured_data,
            }
            figures.append(figure_record)

            evidence_id = f"evidence_{document_id or 'doc'}_p{page_num}_{fig_id}"
            evidence_objects.append({
                "id": evidence_id,
                "documentId": str(document_id),
                "pageNumber": page_num,
                "objectType": "figure",
                "figureType": fig_type,
                "section": current_section,
                "objectId": exact_obj_id,
                "figureNumber": str(fig_num) if fig_num else None,
                "caption": full_caption,
                "imagePath": img_path,
                "content": searchable_content,
                "structuredData": structured_data,
                "bbox": bbox,
                "sourceMethod": "pdf_embedded_image",
                "confidence": 0.95,
                "metadata": {
                    "caption": full_caption,
                    "figureNumber": str(fig_num) if fig_num else None,
                    "objectId": exact_obj_id,
                    "imagePath": img_path,
                    "ocrText": ocr_text,
                    "nearbyContext": nearby_context,
                    "width": width,
                    "height": height,
                },
            })

            sys.stderr.write(f"[FIGURE] document={document_id} page={page_num} figure={fig_id} type={fig_type} status=success\n")

    except Exception as e:
        sys.stderr.write(f"[FIGURE] Page {page_num} figure extraction notice: {e}\n")

    return figures, evidence_objects


class DocumentLoader:
    """
    STAGE 2: Centralized DocumentLoader service.
    Normalizes PDF, DOCX, TXT, CSV, and Images into a unified schema.
    """
    def __init__(self, file_path: str, document_id: str = ""):
        self.path = Path(file_path)
        self.document_id = document_id or self.path.stem
        self.file_name = self.path.name
        self.ext = self.path.suffix.lower()

    def load(self) -> dict:
        if not self.path.exists():
            raise FileNotFoundError(f"File not found: {self.path}")

        if self.ext == ".pdf":
            return self._load_pdf()
        elif self.ext in [".docx", ".doc"]:
            return self._load_docx()
        elif self.ext in [".csv", ".tsv"]:
            return self._load_csv()
        elif self.ext in [".txt", ".md", ".json", ".log", ".xml"]:
            return self._load_text()
        elif self.ext in [".jpg", ".jpeg", ".png", ".bmp", ".tiff"]:
            return self._load_image()
        else:
            return self._load_text()

    def _load_pdf(self) -> dict:
        pages = []
        tables = []
        current_section = "Introduction"
        is_ref = False
        pdf_error = None

        # -------------------------------------------------------------
        # Tier 1: PyMuPDF (fitz) - High performance native text & tables
        # -------------------------------------------------------------
        fitz_module = None
        try:
            import fitz
            fitz_module = fitz
        except ImportError:
            try:
                import pymupdf as fitz
                fitz_module = fitz
            except ImportError:
                fitz_module = None

        if fitz_module is not None:
            try:
                doc = fitz_module.open(str(self.path))
                if doc.is_encrypted or (hasattr(doc, "needs_pass") and doc.needs_pass):
                    doc.close()
                    sys.stderr.write(f"[PDF] Encrypted PDF detected: {self.file_name}\n")
                    return {
                        "success": False,
                        "documentId": self.document_id,
                        "fileName": self.file_name,
                        "fileType": "pdf",
                        "totalPages": 0,
                        "pageCount": 0,
                        "isEncrypted": True,
                        "error": "PDF_PASSWORD_REQUIRED",
                        "pages": [],
                        "tables": [],
                        "fullText": "",
                    }

                total_pages = len(doc)
                sys.stderr.write(f"[PDF] Processing {total_pages} pages via PyMuPDF for {self.file_name}\n")
                all_evidence_objects = []
                figures = []
                tables = []
                text_blocks_count = 0

                for i, page in enumerate(doc, start=1):
                    raw_text = page.get_text("text") or ""
                    cleaned = clean_text_content(raw_text)
                    quality = assess_text_quality(cleaned)
                    method = "native"
                    ocr_used = False
                    ocr_conf = 1.0

                    # Stage 4: OCR Fallback if native text is empty or corrupted
                    if not quality["valid"] or len(cleaned.strip()) < 35:
                        ocr_text, conf = render_and_ocr_page(page, i)
                        ocr_cleaned = clean_text_content(ocr_text)
                        ocr_quality = assess_text_quality(ocr_cleaned)
                        if ocr_quality["valid"] and len(ocr_cleaned) > len(cleaned):
                            cleaned = ocr_cleaned
                            quality = ocr_quality
                            method = "ocr"
                            ocr_used = True
                            ocr_conf = conf

                    current_section, is_ref = detect_section_heading(cleaned, current_section)
                    if i >= total_pages - 2 and re.search(r"\b(references|bibliography)\b", cleaned[:200], re.IGNORECASE):
                        current_section = "References"
                        is_ref = True

                    try:
                        blocks = page.get_text("blocks")
                    except Exception:
                        blocks = []

                    # 1. Multimodal Table Extraction
                    page_tables, table_evidence = extract_tables_from_page(
                        page, i, current_section=current_section, document_id=self.document_id, blocks=blocks
                    )
                    tables.extend(page_tables)
                    all_evidence_objects.extend(table_evidence)

                    # 2. Multimodal Figure & Diagram Extraction
                    page_figures, figure_evidence = extract_figures_from_page(
                        page, i, current_section=current_section, document_id=self.document_id, blocks=blocks, doc=doc
                    )
                    figures.extend(page_figures)
                    all_evidence_objects.extend(figure_evidence)

                    # 3. Layout Text & Heading Blocks
                    page_text_evidence = []
                    for b_idx, b in enumerate(blocks, start=1):
                        b_text = b[4].strip()
                        if not b_text or len(b_text) < 5:
                            continue
                        text_blocks_count += 1
                        
                        if re.match(r"^(?:Table|Tab\.|Figure|Fig\.)\s*\d+", b_text, re.IGNORECASE):
                            continue

                        is_heading = False
                        if re.match(r"^(?:(?:\d+[\.\s]+)*[A-Z][a-zA-Z0-9\s–-]{3,60}|Abstract|Introduction|Methodology|Evaluation|Conclusion|References)$", b_text):
                            is_heading = True
                        elif len(b_text) < 90 and "\n" not in b_text and b_text.isupper():
                            is_heading = True

                        is_equation = bool(re.search(r"(?:Equation|Eq\.)\s*\(?\d+\)?|^\s*[a-zA-Z\d\s=+\-*/^_\\{}()]+\s*=\s*[a-zA-Z\d\s=+\-*/^_\\{}()]+\s*(\(\d+\))?$", b_text))

                        obj_type = "heading" if is_heading else ("equation" if is_equation else "text")
                        obj_id = f"{obj_type}_{i}_{b_idx}"
                        bbox = [b[0], b[1], b[2], b[3]]

                        evidence_item = {
                            "id": f"evidence_{self.document_id}_p{i}_{obj_id}",
                            "documentId": str(self.document_id),
                            "pageNumber": i,
                            "objectType": obj_type,
                            "section": current_section,
                            "objectId": obj_id,
                            "caption": b_text if is_heading else None,
                            "content": f"[{obj_type.upper()} | Page {i} | Section: {current_section}]\n{b_text}",
                            "structuredData": {"text": b_text, "bbox": bbox},
                            "bbox": bbox,
                            "sourceMethod": method if method == "ocr" else "native_text",
                            "confidence": ocr_conf if method == "ocr" else 0.98,
                            "metadata": {"charCount": len(b_text)},
                        }
                        page_text_evidence.append(evidence_item)
                        all_evidence_objects.append(evidence_item)

                    pages.append({
                        "pageNumber": i,
                        "text": cleaned,
                        "tables": page_tables,
                        "figures": page_figures,
                        "evidenceObjects": table_evidence + figure_evidence + page_text_evidence,
                        "section": current_section,
                        "isReference": is_ref,
                        "extractionMethod": method,
                        "ocrUsed": ocr_used,
                        "ocrConfidence": ocr_conf,
                        "qualityScore": quality["score"],
                    })

                doc.close()
            except Exception as e:
                sys.stderr.write(f"[DocumentLoader] PyMuPDF error: {e}\n")
                pdf_error = str(e)

        # -------------------------------------------------------------
        # Tier 2: pdfplumber fallback if PyMuPDF failed or returned 0 pages
        # -------------------------------------------------------------
        if not pages or all(len(p.get("text", "").strip()) == 0 for p in pages):
            try:
                import pdfplumber
                sys.stderr.write(f"[PDF] Falling back to pdfplumber for {self.file_name}\n")
                with pdfplumber.open(str(self.path)) as pdf:
                    pages = []
                    total_pages = len(pdf.pages)
                    for i, page in enumerate(pdf.pages, start=1):
                        extracted_text = page.extract_text() or ""
                        cleaned = clean_text_content(extracted_text)
                        quality = assess_text_quality(cleaned)
                        current_section, is_ref = detect_section_heading(cleaned, current_section)

                        page_tables = []
                        try:
                            t_list = page.extract_tables() or []
                            for t_idx, tab in enumerate(t_list, start=1):
                                if tab and len(tab) >= 2:
                                    headers = [str(c or "").strip() for c in tab[0]]
                                    rows = [[str(c or "").strip() for c in row] for row in tab[1:]]
                                    header_line = "| " + " | ".join(headers) + " |"
                                    sep_line = "| " + " | ".join(["---"] * len(headers)) + " |"
                                    row_lines = ["| " + " | ".join(row) + " |" for row in rows]
                                    raw_text = "\n".join([header_line, sep_line] + row_lines)
                                    tbl_obj = {
                                        "pageNumber": i,
                                        "tableId": f"table_{i}_{t_idx}",
                                        "headers": headers,
                                        "rows": rows,
                                        "rawText": raw_text,
                                        "sourceBBox": [],
                                    }
                                    page_tables.append(tbl_obj)
                                    tables.append(tbl_obj)
                        except Exception as t_err:
                            sys.stderr.write(f"[TABLE] pdfplumber table error on page {i}: {t_err}\n")

                        pages.append({
                            "pageNumber": i,
                            "text": cleaned,
                            "tables": page_tables,
                            "section": current_section,
                            "isReference": is_ref,
                            "extractionMethod": "native_pdfplumber",
                            "ocrUsed": False,
                            "ocrConfidence": 1.0,
                            "qualityScore": quality["score"],
                        })
            except Exception as plumber_err:
                sys.stderr.write(f"[DocumentLoader] pdfplumber fallback error: {plumber_err}\n")
                if not pdf_error:
                    pdf_error = str(plumber_err)

        total_chars = sum(len(p.get("text", "")) for p in pages)
        native_count = sum(1 for p in pages if "native" in p.get("extractionMethod", ""))
        ocr_count = sum(1 for p in pages if p.get("ocrUsed", False))

        sys.stderr.write(f"[PDF] pages={len(pages)} native_text_pages={native_count} ocr_pages={ocr_count} characters={total_chars}\n")

        full_text = "\f".join(p["text"] for p in pages)
        success = len(pages) > 0 and (total_chars > 0 or len(tables) > 0)

        err_msg = None
        if not success:
            if len(pages) > 0 and total_chars == 0:
                # Check if Tesseract binary exists
                t_found = False
                import shutil
                if shutil.which("tesseract"):
                    t_found = True
                else:
                    for t_path in [
                        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                        os.path.expanduser(r"~\AppData\Local\Programs\Tesseract-OCR\tesseract.exe"),
                    ]:
                        if os.path.exists(t_path):
                            t_found = True
                            break
                if not t_found:
                    err_msg = "OCR_UNAVAILABLE"
                else:
                    err_msg = pdf_error or "PDF_EXTRACTION_ZERO_PAGES"
            else:
                err_msg = pdf_error or "PDF_EXTRACTION_ZERO_PAGES"

        diagnostics = {
            "textBlocks": text_blocks_count,
            "tablesDetected": len(tables),
            "tablesParsed": len(tables),
            "figuresDetected": len(figures),
            "figuresProcessed": len(figures),
            "ocrPages": ocr_count,
            "failedObjects": 0,
            "evidenceObjectsCount": len(all_evidence_objects),
        }

        return {
            "success": success,
            "documentId": self.document_id,
            "fileName": self.file_name,
            "fileType": "pdf",
            "pageCount": len(pages),
            "totalPages": len(pages),
            "pages": pages,
            "tables": tables,
            "figures": figures,
            "evidenceObjects": all_evidence_objects,
            "diagnostics": diagnostics,
            "fullText": full_text,
            "error": err_msg if not success else None,
            "errorCode": err_msg if not success else None,
            "metadata": {
                "nativePages": native_count,
                "ocrPages": ocr_count,
                "totalCharacters": total_chars,
                "tablesCount": len(tables),
                "figuresCount": len(figures),
                "evidenceObjectsCount": len(all_evidence_objects),
            },
        }

    def _load_docx(self) -> dict:
        paragraphs = []
        tables = []
        docx_loaded = False
        docx_error = None

        # Strategy A: python-docx if installed
        try:
            import docx
            doc = docx.Document(str(self.path))
            for p in doc.paragraphs:
                txt = p.text.strip()
                if txt:
                    paragraphs.append(txt)

            for t_idx, tab in enumerate(doc.tables, start=1):
                if not tab.rows:
                    continue
                headers = [cell.text.strip() for cell in tab.rows[0].cells]
                rows = []
                for row in tab.rows[1:]:
                    rows.append([cell.text.strip() for cell in row.cells])
                raw_lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
                for r in rows:
                    raw_lines.append("| " + " | ".join(r) + " |")
                table_raw = "\n".join(raw_lines)
                tables.append({
                    "pageNumber": 1,
                    "tableId": f"docx_table_{t_idx}",
                    "headers": headers,
                    "rows": rows,
                    "rawText": table_raw,
                    "sourceBBox": [],
                })
            docx_loaded = True
        except Exception as e:
            sys.stderr.write(f"[DocumentLoader] python-docx unavailable or failed ({e}); using zipfile XML extractor\n")
            docx_error = str(e)

        # Strategy B: Standard library zipfile + WordML XML (100% dependency-free)
        if not docx_loaded or (not paragraphs and not tables):
            try:
                import zipfile
                import xml.etree.ElementTree as ET
                with zipfile.ZipFile(str(self.path)) as z:
                    if "word/document.xml" in z.namelist():
                        xml_data = z.read("word/document.xml")
                        tree = ET.fromstring(xml_data)
                        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                        body = tree.find("w:body", ns)
                        if body is not None:
                            t_idx = 0
                            for child in body:
                                tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
                                if tag == "p":
                                    texts = [node.text for node in child.findall(".//w:t", ns) if node.text]
                                    p_txt = "".join(texts).strip()
                                    if p_txt:
                                        paragraphs.append(p_txt)
                                elif tag == "tbl":
                                    t_idx += 1
                                    rows = []
                                    for tr in child.findall("w:tr", ns):
                                        cells = []
                                        for tc in tr.findall("w:tc", ns):
                                            c_texts = [node.text for node in tc.findall(".//w:t", ns) if node.text]
                                            cells.append("".join(c_texts).strip())
                                        if any(cells):
                                            rows.append(cells)
                                    if rows:
                                        headers = rows[0]
                                        data_rows = rows[1:]
                                        raw_lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
                                        for r in data_rows:
                                            raw_lines.append("| " + " | ".join(r) + " |")
                                        table_raw = "\n".join(raw_lines)
                                        tables.append({
                                            "pageNumber": 1,
                                            "tableId": f"docx_table_{t_idx}",
                                            "headers": headers,
                                            "rows": data_rows,
                                            "rawText": table_raw,
                                            "sourceBBox": [],
                                        })
            except Exception as zip_err:
                sys.stderr.write(f"[DocumentLoader] zipfile WordML extraction error: {zip_err}\n")
                if not docx_error:
                    docx_error = str(zip_err)

        # Combine paragraphs and table text representations
        full_sections = list(paragraphs)
        for tbl in tables:
            raw = tbl.get("rawText")
            if raw and raw not in full_sections:
                full_sections.append(raw)

        text = clean_text_content("\n\n".join(full_sections)) if full_sections else clean_text_content("\n\n".join(paragraphs))
        quality = assess_text_quality(text, file_type="docx")
        section_name, is_ref = detect_section_heading(text, "General")

        pages = [{
            "pageNumber": 1,
            "text": text,
            "tables": tables,
            "section": section_name,
            "isReference": is_ref,
            "extractionMethod": "native",
            "ocrUsed": False,
            "ocrConfidence": 1.0,
            "qualityScore": quality["score"],
        }]

        # Build EvidenceObjects for DOCX
        evidence_objects = []
        for p_idx, p_text in enumerate(paragraphs, start=1):
            is_h = bool(re.match(r"^(?:(?:\d+[\.\s]+)*[A-Z][a-zA-Z0-9\s–-]{3,60}|Abstract|Introduction|Methodology|Evaluation|Conclusion|References)$", p_text))
            obj_type = "heading" if is_h else "text"
            evidence_objects.append({
                "id": f"evidence_{self.document_id}_p1_{obj_type}_{p_idx}",
                "documentId": str(self.document_id),
                "pageNumber": 1,
                "objectType": obj_type,
                "section": section_name,
                "objectId": f"{obj_type}_{p_idx}",
                "caption": p_text if is_h else None,
                "content": f"[{obj_type.upper()} | Page 1 | Section: {section_name}]\n{p_text}",
                "structuredData": {"text": p_text},
                "bbox": [],
                "sourceMethod": "native_text",
                "confidence": 0.98,
                "metadata": {"charCount": len(p_text)},
            })

        for t_idx, tbl in enumerate(tables, start=1):
            t_id = tbl.get("tableId", f"table_{t_idx}")
            t_cap = tbl.get("caption", f"Table {t_idx} on Page 1")
            evidence_objects.append({
                "id": f"evidence_{self.document_id}_p1_{t_id}",
                "documentId": str(self.document_id),
                "pageNumber": 1,
                "objectType": "table",
                "section": section_name,
                "objectId": t_id,
                "caption": t_cap,
                "content": tbl.get("rawText", ""),
                "structuredData": tbl.get("structuredData") or {
                    "tableId": t_id,
                    "caption": t_cap,
                    "columns": tbl.get("headers", []),
                    "rows": tbl.get("rows", []),
                    "rowCount": len(tbl.get("rows", [])),
                    "columnCount": len(tbl.get("headers", [])),
                },
                "bbox": [],
                "sourceMethod": "table_parser",
                "confidence": 0.98,
                "metadata": {"headers": tbl.get("headers", []), "rowCount": len(tbl.get("rows", []))},
            })

        # Extract embedded figures & images from DOCX
        docx_figures = []
        docx_evidence_objects = []
        try:
            import zipfile
            import xml.etree.ElementTree as ET
            from PIL import Image

            script_root = Path(__file__).resolve().parent.parent
            out_dir = script_root / "uploads" / "extracted_images" / str(self.document_id or "doc")
            out_dir.mkdir(parents=True, exist_ok=True)

            with zipfile.ZipFile(str(self.path)) as z:
                rels = {}
                if "word/_rels/document.xml.rels" in z.namelist():
                    r_tree = ET.fromstring(z.read("word/_rels/document.xml.rels"))
                    for r in r_tree:
                        rid = r.attrib.get("Id")
                        target = r.attrib.get("Target")
                        if rid and target:
                            rels[rid] = target

                if "word/document.xml" in z.namelist():
                    doc_tree = ET.fromstring(z.read("word/document.xml"))
                    ns = {
                        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
                        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
                        "v": "urn:schemas-microsoft-com:vml",
                    }
                    body = doc_tree.find("w:body", ns)
                    if body is not None:
                        body_children = list(body)
                        seen_embeds = set()
                        cap_pattern = re.compile(
                            r"^\s*(?:figure|fig\.|chart|diagram|screenshot|illustration)\s*(\d+[a-zA-Z]?|\b[A-Z0-9.-]+)[:.\s]*(.*)$",
                            re.IGNORECASE,
                        )

                        for idx, child in enumerate(body_children):
                            embed_ids = []
                            for blip in child.findall(".//a:blip", ns):
                                embed = blip.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                                if embed:
                                    embed_ids.append(embed)
                            for img_data in child.findall(".//v:imagedata", ns):
                                rid = img_data.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                                if rid:
                                    embed_ids.append(rid)

                            for rid in embed_ids:
                                if rid in seen_embeds:
                                    continue
                                seen_embeds.add(rid)

                                target = rels.get(rid, "")
                                if not target:
                                    continue
                                zip_target = target if target.startswith("word/") else f"word/{target}"
                                if zip_target not in z.namelist():
                                    continue

                                # Search for nearby caption and contextual text
                                caption = ""
                                fig_num = None

                                # Check following paragraphs (up to 4)
                                following_texts = []
                                for next_c in body_children[idx + 1: idx + 5]:
                                    t = "".join([n.text for n in next_c.findall(".//w:t", ns) if n.text]).strip()
                                    if t:
                                        following_texts.append(t)
                                        if not fig_num:
                                            m = cap_pattern.match(t)
                                            if m:
                                                fig_num = m.group(1)
                                                caption = t

                                # Check preceding paragraphs (up to 3)
                                preceding_texts = []
                                for prev_c in body_children[max(0, idx - 3): idx]:
                                    t = "".join([n.text for n in prev_c.findall(".//w:t", ns) if n.text]).strip()
                                    if t:
                                        preceding_texts.append(t)
                                        if not fig_num:
                                            m = cap_pattern.match(t)
                                            if m:
                                                fig_num = m.group(1)
                                                caption = t

                                if not fig_num:
                                    fig_num = str(len(docx_figures) + 1)
                                    caption = f"Figure {fig_num}"

                                # Extract and save image to disk
                                img_bytes = z.read(zip_target)
                                img_filename = f"docx_fig_{fig_num}_{rid}.png"
                                target_path = out_dir / img_filename
                                with open(target_path, "wb") as f_out:
                                    f_out.write(img_bytes)

                                width, height = 0, 0
                                try:
                                    with Image.open(target_path) as im:
                                        width, height = im.size
                                except Exception:
                                    pass

                                # Attempt OCR via pytesseract if available
                                ocr_text = ""
                                try:
                                    import pytesseract
                                    with Image.open(target_path) as im:
                                        ocr_res = pytesseract.image_to_string(im).strip()
                                        if len(ocr_res) > 3:
                                            ocr_text = re.sub(r"\s+", " ", ocr_res)
                                except Exception:
                                    pass

                                nearby_context = " ".join((preceding_texts[-2:] + following_texts[:2]))
                                fig_type = "sequence_diagram" if "sequence" in nearby_context.lower() else ("diagram" if "diagram" in nearby_context.lower() else "figure")
                                exact_obj_id = f"Figure {fig_num}"
                                fig_id = f"figure_{fig_num.lower()}"

                                visual_description = (
                                    f"Figure {fig_num} ({width}x{height} px): {caption}. "
                                    f"Operational context: {nearby_context}"
                                )

                                searchable_content = (
                                    f"[{caption} | Page 1 | Section: {section_name}]\n"
                                    f"Type: {fig_type.replace('_', ' ').capitalize()}\n"
                                    f"Caption: {caption}\n"
                                    f"Visual Description: {visual_description}\n"
                                )
                                if ocr_text:
                                    searchable_content += f"OCR Embedded Text: {ocr_text}\n"
                                if nearby_context:
                                    searchable_content += f"Nearby Text: {nearby_context}\n"

                                structured_data = {
                                    "figureId": fig_id,
                                    "objectId": exact_obj_id,
                                    "figureNumber": str(fig_num),
                                    "page": 1,
                                    "caption": caption,
                                    "figureType": fig_type,
                                    "ocrText": ocr_text,
                                    "visualDescription": visual_description,
                                    "nearbyContext": nearby_context,
                                    "imagePath": str(target_path.resolve()),
                                    "width": width,
                                    "height": height,
                                    "bbox": [],
                                }

                                figure_record = {
                                    "pageNumber": 1,
                                    "figureId": fig_id,
                                    "objectId": exact_obj_id,
                                    "figureNumber": str(fig_num),
                                    "caption": caption,
                                    "figureType": fig_type,
                                    "ocrText": ocr_text,
                                    "visualDescription": visual_description,
                                    "nearbyContext": nearby_context,
                                    "imagePath": str(target_path.resolve()),
                                    "rawText": searchable_content,
                                    "sourceBBox": [],
                                    "structuredData": structured_data,
                                }
                                docx_figures.append(figure_record)

                                evidence_id = f"evidence_{self.document_id}_p1_figure_{fig_num.lower()}"
                                docx_evidence_objects.append({
                                    "id": evidence_id,
                                    "documentId": str(self.document_id),
                                    "pageNumber": 1,
                                    "objectType": "figure",
                                    "figureType": fig_type,
                                    "section": section_name,
                                    "objectId": exact_obj_id,
                                    "figureNumber": str(fig_num),
                                    "caption": caption,
                                    "imagePath": str(target_path.resolve()),
                                    "content": searchable_content,
                                    "structuredData": structured_data,
                                    "bbox": [],
                                    "sourceMethod": "docx_embedded_image",
                                    "confidence": 0.98,
                                    "metadata": {
                                        "caption": caption,
                                        "figureNumber": str(fig_num),
                                        "objectId": exact_obj_id,
                                        "imagePath": str(target_path.resolve()),
                                        "ocrText": ocr_text,
                                        "nearbyContext": nearby_context,
                                        "width": width,
                                        "height": height,
                                    },
                                })

                                sys.stderr.write(f"[DOCX_IMAGE] document={self.document_id} relation={rid} image={zip_target} status=extracted\n")
                                sys.stderr.write(f"[FIGURE] document={self.document_id} figure={fig_num} caption_found=true ocr={bool(ocr_text)} vision=pending status=indexed\n")

        except Exception as docx_img_err:
            sys.stderr.write(f"[DOCX_IMAGE_ERR] Failed extracting images from DOCX: {docx_img_err}\n")

        evidence_objects.extend(docx_evidence_objects)
        pages[0]["evidenceObjects"] = evidence_objects
        pages[0]["figures"] = docx_figures

        total_chars = len(text)
        success = bool(paragraphs or tables or docx_figures)

        diagnostics = {
            "textBlocks": len(paragraphs),
            "tablesDetected": len(tables),
            "tablesParsed": len(tables),
            "figuresDetected": len(docx_figures),
            "figuresProcessed": len(docx_figures),
            "ocrPages": 0,
            "failedObjects": 0,
            "evidenceObjectsCount": len(evidence_objects),
        }

        return {
            "success": success,
            "documentId": self.document_id,
            "fileName": self.file_name,
            "fileType": "docx",
            "pageCount": 1,
            "totalPages": 1,
            "pages": pages,
            "tables": tables,
            "figures": docx_figures,
            "evidenceObjects": evidence_objects,
            "diagnostics": diagnostics,
            "fullText": text,
            "error": None if success else (docx_error or "DOCX_EXTRACTION_ZERO_CONTENT"),
            "metadata": {
                "paragraphsCount": len(paragraphs),
                "tablesCount": len(tables),
                "figuresCount": len(docx_figures),
                "evidenceObjectsCount": len(evidence_objects),
                "totalCharacters": total_chars,
            },
        }

    def _load_csv(self) -> dict:
        rows = []
        delimiter = "\t" if self.ext == ".tsv" else ","
        try:
            with open(self.path, mode="r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f, delimiter=delimiter)
                for r in reader:
                    if r:
                        rows.append(r)
        except Exception as e:
            sys.stderr.write(f"[DocumentLoader] CSV read error: {e}\n")

        headers = rows[0] if rows else []
        data_rows = rows[1:] if len(rows) > 1 else []
        raw_lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
        for r in data_rows[:200]:
            raw_lines.append("| " + " | ".join(r) + " |")

        table_md = "\n".join(raw_lines)
        table_obj = {
            "pageNumber": 1,
            "tableId": "csv_table_1",
            "headers": headers,
            "rows": data_rows,
            "rawText": table_md,
            "sourceBBox": [],
        }
        text_rep = f"CSV Dataset: {self.file_name}\n\n{table_md}"
        return {
            "success": True,
            "documentId": self.document_id,
            "fileName": self.file_name,
            "fileType": "csv",
            "pageCount": 1,
            "totalPages": 1,
            "pages": [{
                "pageNumber": 1,
                "text": text_rep,
                "tables": [table_obj],
                "section": "Dataset",
                "isReference": False,
                "extractionMethod": "native",
                "ocrUsed": False,
                "ocrConfidence": 1.0,
                "qualityScore": 0.95,
            }],
            "tables": [table_obj],
            "fullText": table_md,
            "metadata": {
                "rowCount": len(rows),
                "colCount": len(headers),
            },
        }

    def _load_text(self) -> dict:
        raw = self.path.read_text(encoding="utf-8", errors="ignore")
        cleaned = clean_text_content(raw)
        quality = assess_text_quality(cleaned)
        return {
            "success": len(cleaned) > 0,
            "documentId": self.document_id,
            "fileName": self.file_name,
            "fileType": self.ext.replace(".", ""),
            "pageCount": 1,
            "totalPages": 1,
            "pages": [{
                "pageNumber": 1,
                "text": cleaned,
                "tables": [],
                "section": "General",
                "isReference": False,
                "extractionMethod": "native",
                "ocrUsed": False,
                "ocrConfidence": 1.0,
                "qualityScore": quality["score"],
            }],
            "tables": [],
            "fullText": cleaned,
            "metadata": {
                "totalCharacters": len(cleaned),
            },
        }

    def _load_image(self) -> dict:
        ocr_text = ""
        conf = 0.0
        try:
            from PIL import Image
            import pytesseract
            for t_path in [
                r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
                os.path.expanduser(r"~\AppData\Local\Programs\Tesseract-OCR\tesseract.exe"),
            ]:
                if os.path.exists(t_path):
                    pytesseract.pytesseract.tesseract_cmd = t_path
                    break
            img = Image.open(self.path)
            ocr_text = pytesseract.image_to_string(img).strip()
            conf = 0.85
        except Exception as e:
            sys.stderr.write(f"[DocumentLoader] Image OCR error: {e}\n")

        cleaned = clean_text_content(ocr_text)
        quality = assess_text_quality(cleaned)
        return {
            "success": len(cleaned) > 0,
            "documentId": self.document_id,
            "fileName": self.file_name,
            "fileType": "image",
            "pageCount": 1,
            "totalPages": 1,
            "pages": [{
                "pageNumber": 1,
                "text": cleaned,
                "tables": [],
                "section": "Scanned Document",
                "isReference": False,
                "extractionMethod": "ocr",
                "ocrUsed": True,
                "ocrConfidence": conf,
                "qualityScore": quality["score"],
            }],
            "tables": [],
            "fullText": cleaned,
            "metadata": {
                "ocrConfidence": conf,
                "totalCharacters": len(cleaned),
            },
        }


def main():
    if len(sys.argv) < 2:
        sys.stderr.write("Usage: document_loader.py <file_path> [--json]\n")
        sys.exit(1)

    file_path = sys.argv[1]
    loader = DocumentLoader(file_path)
    result = loader.load()

    if "--json" in sys.argv:
        sys.stdout.write(json.dumps(result, ensure_ascii=False) + "\n")
        sys.stdout.flush()
    else:
        sys.stdout.buffer.write(result.get("fullText", "").encode("utf-8", errors="replace"))
        sys.stdout.buffer.flush()


if __name__ == "__main__":
    main()
